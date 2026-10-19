"""
ponte.py — leva os movimentos aprovados na PWA para o Organic, sem nunca escrever no Excel.

Uso:
  python ponte.py                 importa: Sheet (aprovados na PWA E SMS ainda por aprovar) -> lote JSON -> organic.py propor (folha Pendentes)
                                  o Rei só decide no Organic: aprovar / corrigir / regra
  python ponte.py --seco          mostra o lote que seria criado; não escreve nada
  python ponte.py --saldos        só copia os saldos de Contas para a Sheet (a PWA mostra-os)
  python ponte.py --revisao       revisão semanal: últimos 7 dias do Diario, pendentes, saldos; actualiza _Estado.md
  python ponte.py --conferir      confere o saldo M-Pesa dos SMS com o saldo calculado pelos movimentos da Sheet
  python ponte.py --de ficheiro   usa movimentos de um JSON local (testes); não toca na Sheet

Requer: pip install requests openpyxl
Lê GAS_URL e GAS_TOKEN do .env desta pasta. ORGANIC_DIR vem do .env ou do ambiente; se faltar, usa a pasta THE ORANIC do Rei.
O registo oficial é Contabilidade.xlsx: só scripts/organic.py do Organic escreve nele (fecha o Excel antes).
"""
import json
import os
import re
import shutil
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import load_workbook

AQUI = Path(__file__).parent
ORGANIC_POR_DEFEITO = r"C:\Users\Edson Adolfo\Desktop\AI\The castelo do Rei multiplicado\THE ORANIC"

# Contas da PWA -> contas do Organic (Caixa e M-Pesa estão unidos em CASH)
CONTAS = {"Caixa": "CASH", "M-Pesa": "CASH", "e-Mola": "CASH", "Banco": "BANCO"}

# Categorias da PWA -> código do Plano_Contas. Patrimoniais tratadas à parte (FIXER).
CAT_SAIDA = {
    "Alimentação": "D-ALIM", "Transporte": "D-TRANS", "Casa": "D-CASA", "Comunicação": "D-COM",
    "Saúde": "D-SAUDE", "Pessoal": "D-OUT", "Lazer": "D-LAZER", "Educação": "D-EDUC",
    "Dízimo": "D-DIZ", "Família": "D-FAM", "Software": "D-SOFT", "Equipamento": "D-EQUIP",
    "Taxas": "D-BANC", "Outros": "D-OUT",
}
CAT_ENTRADA = {"Rendimento": "R-OUT", "Salário": "R-SAL", "Outros": "R-OUT"}
FIXER_SAIDA = "Aporte Fixer"        # -> I-APORTE (alternativa E-CONC: decide o Rei no Organic)
FIXER_ENTRADA = "Reembolso Fixer"   # -> E-REEMB (alternativa I-RESGATE)
IGNORAR = "Saldo inicial"           # os saldos iniciais já estão na folha Contas


def env():
    cfg = {}
    f = AQUI / ".env"
    if f.exists():
        for linha in f.read_text(encoding="utf-8").splitlines():
            linha = linha.strip()
            if linha and not linha.startswith("#") and "=" in linha:
                k, v = linha.split("=", 1)
                cfg[k.strip()] = v.strip()
    for k in ("GAS_URL", "GAS_TOKEN", "ORGANIC_DIR"):
        if os.environ.get(k):
            cfg[k] = os.environ[k]
    return cfg


def gas(cfg, corpo=None, **params):
    """Chama o GAS. O Google devolve 404/HTML de vez em quando: repete até 4 vezes."""
    for i in range(4):
        try:
            if corpo is None:
                r = requests.get(cfg["GAS_URL"], params={"token": cfg["GAS_TOKEN"], **params}, timeout=60)
            else:
                r = requests.post(cfg["GAS_URL"], data=json.dumps({"token": cfg["GAS_TOKEN"], **corpo}), timeout=60)
            j = r.json()
            if j.get("erro") == "token inválido" and i < 3:  # o GAS falha isto de vez em quando
                raise ValueError("token")
            break
        except ValueError:
            if i == 3:
                sys.exit("O GAS não devolveu JSON (HTTP %s). Tenta outra vez daqui a pouco." % r.status_code)
            time.sleep(2)
    if not j.get("ok"):
        sys.exit("Erro do GAS: " + str(j.get("erro")))
    return j


def e_std(m):
    """STD Bank é conta da Fixer (STD Fixer 460): fica fora do Organic. 'Banco' com SMS do Millennium BIM ou registo manual é pessoal."""
    if m.get("conta") == "Banco Standard":
        return True
    if m.get("conta") == "Banco" and m.get("sms"):
        return not re.search(r"MILLENNIUM|MILLENIUM|BIM", m["sms"].upper())
    return False


def conta(c):
    if c not in CONTAS:
        raise ValueError(f"conta desconhecida na PWA: {c!r}")
    return CONTAS[c]


def para_lote(m):
    """Converte um movimento da PWA num movimento do lote do Organic, com a confiança ajustada ao estado."""
    d, aviso = _para_lote(m)
    if d and m.get("estado") == "pendente":
        generico = (m.get("descricao") or "") == "SMS por classificar" or not m.get("categoria")
        d["confianca"] = "baixa" if generico else "media"
        d["nota"] = (d.get("nota", "") + " SMS sem aprovação na PWA: confirmar no Organic").strip()
    return d, aviso


def _para_lote(m):
    """Converte um movimento da PWA num movimento do lote do Organic. Devolve (dict|None, aviso|None)."""
    if not float(m.get("valor") or 0):
        return None, f"{m['id']}: valor 0 ({m.get('descricao')}) — ignorado"
    base = {"data": m["data"], "referencia": m["id"],
            "fonte": "caixa-sms" if m.get("origem") == "sms" else "caixa-pwa", "confianca": "alta"}
    if e_std(m):
        return None, f"{m['id']}: movimento do STD Bank (conta da Fixer) fora do Organic — {m['valor']:.2f} MT, {m['data']}"
    desc = m.get("descricao") or m.get("categoria") or ""
    v = float(m["valor"])
    cat = m.get("categoria", "")
    if m["tipo"] == "transferencia":
        de, para = conta(m["conta"]), conta(m["conta_destino"])
        if de == para:
            return None, f"{m['id']}: transferência entre contas unidas ({m['conta']} -> {m['conta_destino']}) não precisa de registo"
        return {**base, "categoria": "T-INT", "valor": v, "descricao": desc or "Transferência",
                "conta_saida": de, "conta_entrada": para}, None
    c = conta(m["conta"])
    if m["tipo"] == "saida":
        if cat == FIXER_SAIDA:
            return {**base, "categoria": "I-APORTE", "valor": v, "descricao": desc or "Aporte Fixer",
                    "conta_saida": c, "conta_entrada": "INVEST_FIXER", "confianca": "baixa",
                    "nota": "Recomendado: aporte de capital (I-APORTE). Se for empréstimo, corrigir para E-CONC e conta_entrada=EMPREST_FIXER"}, None
        cod = CAT_SAIDA.get(cat)
        if cod is None:
            base["confianca"] = "baixa"
        return {**base, "categoria": cod or "D-OUT", "saida": v, "desc_saida": desc, "conta_saida": c,
                **({"nota": f"Categoria da PWA sem equivalente: {cat!r}"} if cod is None else {})}, None
    if cat == IGNORAR:
        return None, f"{m['id']}: 'Saldo inicial' ignorado (os saldos já estão em Contas)"
    if cat == FIXER_ENTRADA:
        return {**base, "categoria": "E-REEMB", "valor": v, "descricao": desc or "Reembolso Fixer",
                "conta_entrada": c, "conta_saida": "EMPREST_FIXER", "confianca": "baixa",
                "nota": "Recomendado: reembolso de empréstimo (E-REEMB). Se for resgate de capital, corrigir para I-RESGATE e conta_saida=INVEST_FIXER"}, None
    cod = CAT_ENTRADA.get(cat)
    if cod is None:
        base["confianca"] = "baixa"
    return {**base, "categoria": cod or "R-OUT", "entrada": v, "desc_entrada": desc, "conta_entrada": c,
            **({"nota": f"Categoria da PWA sem equivalente: {cat!r}"} if cod is None else {})}, None


def caminhos(cfg):
    base = Path(cfg.get("ORGANIC_DIR") or ORGANIC_POR_DEFEITO)
    xlsx = base / "01_OPERACIONAL" / "Contabilidade" / "Contabilidade.xlsx"
    script = base / "scripts" / "organic.py"
    if not xlsx.exists() or not script.exists():
        sys.exit(f"ORGANIC_DIR inválido (falta Contabilidade.xlsx ou scripts/organic.py): {base}")
    return base, xlsx, script


def saldos(xlsx):
    """Saldo de cada conta = saldo inicial + entradas - saídas do Diario (só leitura)."""
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ct = list(wb["Contas"].iter_rows(values_only=True))
    h = [str(x or "").strip().lower() for x in ct[0]]
    i_c, i_s = h.index("conta"), next(i for i, x in enumerate(h) if x.startswith("saldo inicial"))
    s = {r[i_c]: float(r[i_s] or 0) for r in ct[1:] if r[i_c]}
    dr = list(wb["Diario"].iter_rows(values_only=True))
    dh = [str(x or "").strip().lower() for x in dr[0]]
    ie, ce = dh.index("entrada (mzn)"), dh.index("conta entrada")
    isd, cs = dh.index("saída (mzn)"), dh.index("conta saída")
    for r in dr[1:]:
        if r[ce] in s and r[ie]:
            s[r[ce]] += float(r[ie])
        if r[cs] in s and r[isd]:
            s[r[cs]] -= float(r[isd])
    wb.close()
    return {k: round(v, 2) for k, v in s.items()}


RE_SALDO = re.compile(r"saldo\s+M-?Pesa\s+e\s+de\s+([\d.,]+)\s*MT", re.I)


def num(t):
    t = t.rstrip(".,")
    if "," in t and "." in t:
        t = t.replace(",", "") if t.rfind(".") > t.rfind(",") else t.replace(".", "").replace(",", ".")
    elif "," in t:
        t = t.replace(",", ".") if re.search(r",\d{2}$", t) else t.replace(",", "")
    return float(t)


def efeito_mpesa(m):
    """Quanto o movimento mexe no saldo M-Pesa (a taxa é um movimento à parte)."""
    v = float(m["valor"])
    if m["tipo"] == "transferencia":
        return (-v if m["conta"] == "M-Pesa" else 0) + (v if m.get("conta_destino") == "M-Pesa" else 0)
    if m["conta"] != "M-Pesa":
        return 0
    return v if m["tipo"] == "entrada" else -v


RE_QUANDO = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{2,4})\s+as\s+(\d{1,2}):(\d{2})\s*(AM|PM)", re.I)


def quando(m):
    """Hora real do SMS M-Pesa ('aos 3/10/26 as 11:20 AM'); sem ela, a data de criação. A ordem da Sheet não é a do tempo."""
    q = RE_QUANDO.search(m.get("sms") or "")
    if not q:
        return (m.get("criado_em") or m["data"]).replace("T", " ")
    d, mes, a, h, mi, ap = int(q[1]), int(q[2]), int(q[3]), int(q[4]), int(q[5]), q[6].upper()
    h = h % 12 + (12 if ap == "PM" else 0)
    return f"{a + 2000 if a < 100 else a:04d}-{mes:02d}-{d:02d} {h:02d}:{mi:02d}"


def todos_movimentos(cfg):
    """Todos os movimentos (acção 'todos' do GAS). Se o GAS ainda for antigo: pendentes + recentes + por importar (incompleto)."""
    j = gas(cfg, acao="todos")
    if "movimentos" in j:
        return j["movimentos"]
    print("  (GAS antigo: sem a acção 'todos', a conferência pode dar diferenças falsas até haver nova versão)")
    vistos = {}
    for m in gas(cfg, acao="listar")["pendentes"] + gas(cfg, acao="listar")["recentes"] + gas(cfg, acao="exportar")["movimentos"]:
        vistos[m["id"]] = m
    return list(vistos.values())


def conferir(movs, auto=False):
    """Segue o saldo M-Pesa SMS a SMS: saldo anterior + movimentos tem de dar o saldo do SMS seguinte.
    Uma diferença = falta um movimento (SMS não lido, levantamento sem SMS, etc.). Devolve o nº de diferenças."""
    grupos = {}
    for m in sorted((x for x in movs if x.get("estado") != "rejeitado"), key=quando):
        if efeito_mpesa(m) == 0:
            continue
        grupos.setdefault(re.sub(r"-tx$", "", str(m["id"])), []).append(m)
    corrente, ancora, difs, ultimo_sms, flut = None, None, [], None, {}
    for gid, ms in grupos.items():
        ef = round(sum(efeito_mpesa(m) for m in ms), 2)
        sms = next((x for m in ms for x in [RE_SALDO.search(m.get("sms") or "")] if x), None)
        if not sms and not any(RE_QUANDO.search(m.get("sms") or "") for m in ms):
            flut[gid] = ef  # sem saldo nem hora (ex. recibo Credelec, registo manual): não se sabe onde encaixa na cadeia
            continue
        if corrente is None:
            if sms:
                corrente = ancora = num(sms.group(1))
                ultimo_sms = (corrente, ms[0]["data"])
            continue
        corrente = round(corrente + ef, 2)
        if sms:
            lido = num(sms.group(1))
            dif = round(lido - corrente, 2)
            if abs(dif) > 0.009:
                expl = next((g for g, e in flut.items() if abs(e - dif) < 0.009), None)
                if expl:
                    flut.pop(expl)
                else:
                    difs.append((ms[0]["data"], gid, lido, corrente, dif))
            corrente, ultimo_sms = lido, (lido, ms[0]["data"])
    if ancora is None:
        if not auto:
            print("Conferência M-Pesa: ainda não há SMS com saldo.")
        return 0
    print(f"Conferência M-Pesa — último saldo nos SMS: {ultimo_sms[0]:,.2f} MT ({ultimo_sms[1]}); saldo pelos movimentos: {corrente:,.2f} MT")
    for d, gid, lido, calc, dif in difs:
        sinal = "entrou" if dif > 0 else "saiu"
        print(f"  ! {d} {gid}: SMS diz {lido:,.2f}, movimentos dão {calc:,.2f} -> {abs(dif):,.2f} MT {sinal} sem registo antes deste SMS")
    if abs(corrente - ultimo_sms[0]) > 0.009:
        print(f"  ! Depois do último SMS há movimentos manuais M-Pesa que o mexem em {corrente - ultimo_sms[0]:+,.2f} MT")
    if flut:
        print("  · sem hora no SMS (não entram na cadeia): " + ", ".join(f"{g} {e:+,.2f}" for g, e in flut.items()))
    if not difs and abs(corrente - ultimo_sms[0]) <= 0.009:
        print("  OK — os movimentos explicam todos os saldos dos SMS.")
    return len(difs)


def revisao(cfg, xlsx, script):
    """Resumo dos últimos 7 dias (só leitura do Diario) + pendentes + saldos."""
    from collections import defaultdict
    from datetime import date, timedelta
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    rows = list(wb["Diario"].iter_rows(values_only=True))
    h = [str(x or "").strip().lower() for x in rows[0]]
    i_d, i_e, i_s, i_c = h.index("data"), h.index("entrada (mzn)"), h.index("saída (mzn)"), h.index("categoria")
    pend = sum(1 for r in wb["Pendentes"].iter_rows(min_row=2, values_only=True) if any(v not in (None, "") for v in r[1:]))
    wb.close()
    desde = date.today() - timedelta(days=6)
    ent, sai = defaultdict(float), defaultdict(float)
    for r in rows[1:]:
        d = r[i_d]
        if not d or not hasattr(d, "date") or d.date() < desde:
            continue
        cat = str(r[i_c] or "?")
        if cat.startswith("R-"):
            ent[cat] += float(r[i_e] or 0)
        elif cat.startswith("D-"):
            sai[cat] += float(r[i_s] or 0)
    print(f"Revisão {desde:%d-%m} a {date.today():%d-%m}")
    print(f"  Receitas {sum(ent.values()):,.2f}  ·  Despesas {sum(sai.values()):,.2f}  ·  Resultado {sum(ent.values()) - sum(sai.values()):,.2f} MZN")
    for cat, v in sorted(sai.items(), key=lambda x: -x[1]):
        print(f"    {cat:<10}{v:>12,.2f}")
    print(f"  Pendentes no Organic: {pend}")
    try:
        j = gas(cfg, acao="listar")
        print(f"  Na PWA: {len(j['pendentes'])} por aprovar, {j['por_importar']} aprovados por passar ao Organic")
    except SystemExit as e:
        print("  PWA indisponível:", e)
    print("  Saldos:", ", ".join(f"{k} {v:,.2f}" for k, v in saldos(xlsx).items() if v) or "tudo a zero")
    try:
        conferir(todos_movimentos(cfg))
    except SystemExit as e:
        print("  Conferência M-Pesa indisponível:", e)
    subprocess.run([sys.executable, str(script), "estado"], cwd=str(script.parent.parent))


def main(a):
    cfg = env()
    seco, so_saldos = "--seco" in a, "--saldos" in a
    de = a[a.index("--de") + 1] if "--de" in a else None
    base, xlsx, script = caminhos(cfg)
    if "--revisao" in a:
        return revisao(cfg, xlsx, script)
    if "--conferir" in a:
        return conferir(todos_movimentos(cfg))

    if not so_saldos:
        if de:
            movs = json.loads(Path(de).read_text(encoding="utf-8"))
        else:
            movs = gas(cfg, acao="exportar")["movimentos"]
            ids = {m["id"] for m in movs}
            movs += [m for m in gas(cfg, acao="listar")["pendentes"] if m["id"] not in ids]  # SMS por aprovar: vão já para o Organic
        movs.sort(key=lambda m: (m["data"], m.get("criado_em", "")))
        lote, feitos = [], []
        for m in movs:
            try:
                d, aviso = para_lote(m)
            except ValueError as e:
                print(f"! {m['id']}: {e} — deixado por importar")
                continue
            if aviso:
                print("· " + aviso)
            if d:
                lote.append(d)
            feitos.append(m["id"])
        print(f"{len(movs)} aprovado(s) na PWA -> {len(lote)} para o Organic")
        if not de:
            try:
                conferir(todos_movimentos(cfg), auto=True)
            except SystemExit as e:
                print("Conferência M-Pesa indisponível:", e)
        if seco:
            print(json.dumps(lote, ensure_ascii=False, indent=1))
            return
        if not lote:
            if feitos and not de:
                gas(cfg, {"acao": "importado", "ids": feitos})
            return
        agora = datetime.now().strftime("%Y%m%d-%H%M%S")
        pasta = base / "00_ENTRADA" / "_lotes"
        pasta.mkdir(parents=True, exist_ok=True)
        ficheiro = pasta / f"caixa-{agora}.json"
        ficheiro.write_text(json.dumps(lote, ensure_ascii=False, indent=1), encoding="utf-8")
        bk = xlsx.parent / "backup"
        bk.mkdir(exist_ok=True)
        shutil.copy2(xlsx, bk / f"Contabilidade-{agora}.xlsx")
        r = subprocess.run([sys.executable, str(script), "propor", str(ficheiro)], cwd=str(base))
        if r.returncode != 0:
            sys.exit("organic.py propor falhou (o Excel está aberto?). Nada marcado como importado; corre outra vez.")
        if not de:
            gas(cfg, {"acao": "importado", "ids": feitos})
        print(f"Lote {ficheiro.name} proposto. Aprova em Pendentes: python scripts/organic.py listar / aprovar")

    if so_saldos or not de:
        s = saldos(xlsx)
        gas(cfg, {"acao": "saldos", "saldos": s})
        print("Saldos de Contas enviados para a Sheet:", ", ".join(f"{k} {v:,.2f}" for k, v in s.items() if v))


if __name__ == "__main__":
    main(sys.argv[1:])
