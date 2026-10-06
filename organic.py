"""
organic.py — sincroniza os movimentos aprovados (Google Sheet via GAS) para o Excel do Organic.

Uso:   python organic.py sincronizar
Requer: pip install openpyxl requests
Lê GAS_URL, GAS_TOKEN e EXCEL_PATH do ficheiro .env nesta pasta.
O Excel é regenerado a partir da Sheet (a Sheet é a fonte de verdade).
"""
import sys
from collections import defaultdict
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

AQUI = Path(__file__).parent


def env():
    cfg = {}
    for linha in (AQUI / ".env").read_text(encoding="utf-8").splitlines():
        linha = linha.strip()
        if linha and not linha.startswith("#") and "=" in linha:
            k, v = linha.split("=", 1)
            cfg[k.strip()] = v.strip()
    return cfg


def buscar(cfg):
    r = requests.get(cfg["GAS_URL"], params={"acao": "exportar", "token": cfg["GAS_TOKEN"]}, timeout=60)
    j = r.json()
    if not j.get("ok"):
        sys.exit("Erro do GAS: " + str(j.get("erro")))
    return sorted(j["movimentos"], key=lambda m: (m["data"], m.get("criado_em", "")))


def efeito(m):
    """Devolve [(conta, delta)] do movimento."""
    v = float(m["valor"])
    if m["tipo"] == "entrada":
        return [(m["conta"], v)]
    if m["tipo"] == "saida":
        return [(m["conta"], -v)]
    return [(m["conta"], -v), (m["conta_destino"], v)]


CAB = Font(bold=True, color="FFFFFF")
FUNDO = PatternFill("solid", fgColor="1F3A28")
MZN = '#,##0.00'


def cabecalho(ws, cols, larguras):
    ws.append(cols)
    for i, c in enumerate(ws[1], 1):
        c.font, c.fill, c.alignment = CAB, FUNDO, Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = larguras[i - 1]
    ws.freeze_panes = "A2"


def gerar(movs, caminho):
    wb = Workbook()

    # Diário geral (como no caderno: com saldo corrido)
    d = wb.active
    d.title = "Diario"
    cabecalho(d, ["Data", "Tipo", "Conta", "Destino", "Categoria", "Descrição", "Entrada", "Saída", "Saldo total", "Origem", "ID"],
              [12, 14, 12, 12, 16, 32, 13, 13, 14, 9, 11])
    total = 0.0
    for m in movs:
        v = float(m["valor"])
        ent = v if m["tipo"] == "entrada" else None
        sai = v if m["tipo"] == "saida" else None
        total += (ent or 0) - (sai or 0)
        d.append([m["data"], m["tipo"], m["conta"], m.get("conta_destino") or "", m.get("categoria") or "",
                  m.get("descricao") or "", ent, sai, total, m.get("origem"), m["id"]])
    for col in "GHI":
        for c in d[col][1:]:
            c.number_format = MZN

    # Saldos por conta
    s = wb.create_sheet("Saldos")
    cabecalho(s, ["Conta", "Saldo (MZN)"], [16, 16])
    saldos = defaultdict(float)
    for m in movs:
        for conta, delta in efeito(m):
            saldos[conta] += delta
    for conta, v in sorted(saldos.items()):
        s.append([conta, v])
    s.append(["TOTAL", sum(saldos.values())])
    s.cell(s.max_row, 1).font = Font(bold=True)
    for c in s["B"][1:]:
        c.number_format = MZN

    # Resumo mensal por categoria (saídas e entradas, sem transferências)
    r = wb.create_sheet("Mensal")
    cabecalho(r, ["Mês", "Tipo", "Categoria", "Valor (MZN)"], [10, 10, 18, 14])
    agg = defaultdict(float)
    for m in movs:
        if m["tipo"] != "transferencia":
            agg[(m["data"][:7], m["tipo"], m.get("categoria") or "Outros")] += float(m["valor"])
    for (mes, tipo, cat), v in sorted(agg.items()):
        r.append([mes, tipo, cat, v])
    for c in r["D"][1:]:
        c.number_format = MZN

    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
    return saldos


def main():
    if len(sys.argv) < 2 or sys.argv[1] != "sincronizar":
        sys.exit("Uso: python organic.py sincronizar")
    cfg = env()
    movs = buscar(cfg)
    caminho = Path(cfg["EXCEL_PATH"])
    saldos = gerar(movs, caminho)
    print(f"{len(movs)} movimentos → {caminho}")
    for conta, v in sorted(saldos.items()):
        print(f"  {conta:<10} {v:>14,.2f} MZN")


if __name__ == "__main__":
    main()
